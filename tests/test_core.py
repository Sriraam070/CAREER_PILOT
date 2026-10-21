import json

import pytest

from careerpilot import api, apply, db, discover, evaluate, inbox, reports
from careerpilot import sources


def job(**kw):
    base = {"company": "ACME", "title": "Python Developer", "location": "Chennai, India", "url": "https://x/1",
            "description": "Requirements:\n- 2+ years Python, Django, PostgreSQL, Docker, AWS\nNice to have:\n- Kubernetes",
            "source": "test", "posted_at": None}
    base.update(kw)
    return base


def test_dedup_across_sources():
    a, new_a = db.upsert_job(job(source="workday", url="https://a"))
    b, new_b = db.upsert_job(job(source="adzuna", url="https://b", title="Python  Developer!", location="Chennai"))
    assert new_a and not new_b and a == b
    c, new_c = db.upsert_job(job(title="Java Developer", url="https://c"))
    assert new_c and c != a


@pytest.mark.parametrize("label", ["Current CTC (LPA)", "Expected Salary", "Notice period in days", "I agree to the terms",
                                   "Electronic signature", "Password", "Aadhaar number", "Date of Birth", "Gender",
                                   "Do you require visa sponsorship?", "I hereby declare that the information is true"])
def test_sensitive_labels(label):
    assert apply.is_sensitive(label)


@pytest.mark.parametrize("label", ["First Name", "Email", "Phone", "LinkedIn Profile", "GitHub URL", "Resume/CV"])
def test_safe_labels(label):
    assert not apply.is_sensitive(label)


def test_plan_fields_never_fills_sensitive():
    fields = [{"idx": 0, "tag": "input", "type": "text", "label": "First Name *", "required": True, "name": "fn"},
              {"idx": 1, "tag": "input", "type": "email", "label": "Email", "required": True, "name": "em"},
              {"idx": 2, "tag": "input", "type": "text", "label": "Current CTC", "required": True, "name": "ctc"},
              {"idx": 3, "tag": "input", "type": "checkbox", "label": "I consent to processing", "required": True, "name": "c"},
              {"idx": 4, "tag": "input", "type": "text", "label": "How did you hear about us?", "required": True, "name": "h"},
              {"idx": 5, "tag": "input", "type": "file", "label": "Resume", "required": True, "name": "r"},
              {"idx": 6, "tag": "input", "type": "password", "label": "Create account", "required": True, "name": "p"}]
    fill, sensitive, unknown, resume = apply.plan_fields(fields, {"first_name": "Test", "email": "t@e.com"})
    assert [f[0] for f in fill] == [0, 1]
    assert "Current CTC" in sensitive and "I consent to processing" in sensitive and "Create account" in sensitive
    assert unknown == ["How did you hear about us?"]
    assert resume["idx"] == 5


def test_evaluate_and_probability_is_bounded():
    jid, _ = db.upsert_job(job())
    r = evaluate.evaluate(jid)
    assert 0 < r["match"] <= 100 and "PostgreSQL" in r["gaps"]
    assert 1 <= r["probability"] <= 60
    p_low, _ = evaluate.probability(20, 8, 2.5, ["MARQUEE"], 60)
    p_high, _ = evaluate.probability(90, 2, 2.5, ["DAY SHARING"], 1)
    assert p_low < p_high <= 60


def test_preferences_filter():
    prefs = {"exclude_keywords": ["senior"], "locations": ["Chennai"], "max_years_required": 4, "min_salary_lpa": 8}
    assert evaluate.passes_preferences(job(), prefs)[0]
    assert not evaluate.passes_preferences(job(title="Senior Python Developer"), prefs)[0]
    assert not evaluate.passes_preferences(job(location="Berlin"), prefs)[0]
    assert not evaluate.passes_preferences({**job(), "years_required": 7}, prefs)[0]
    assert not evaluate.passes_preferences(job(salary_text="₹4L - ₹6L"), prefs)[0]


def test_salary_parse():
    assert evaluate.parse_salary_lpa("CTC 6 - 9 LPA") == (6.0, 9.0)
    assert evaluate.parse_salary_lpa("₹12L") == (12.0, 12.0)
    assert evaluate.parse_salary_lpa("competitive") is None


@pytest.mark.parametrize("subject,body,expected", [
    ("Thank you for applying to Zoho", "We have received your application", "Applied"),
    ("Your HackerRank test", "Complete the online test by Friday", "OA/Test"),
    ("Interview invitation", "We'd like to invite you for an interview", "Interview"),
    ("Update on your application", "Unfortunately we will not be moving forward", "Rejected"),
    ("Offer", "We are pleased to offer you the role", "Offer"),
    ("Newsletter", "Our product update", None)])
def test_inbox_classify(subject, body, expected):
    assert inbox.classify(subject, body) == expected


def test_adapters_parse(monkeypatch):
    from careerpilot.sources import amazon, ashby, greenhouse, lever, smartrecruiters, workday
    responses = {
        "workday_list": {"jobPostings": [{"title": "Software Engineer - Python", "externalPath": "/job/Chennai/SE_R1",
                                          "locationsText": "Chennai", "postedOn": "Posted 2 Days Ago", "bulletFields": ["R1"]},
                                         {"title": "Marketing Manager", "externalPath": "/job/x/MM", "locationsText": "Pune"}]},
        "workday_detail": {"jobPostingInfo": {"jobDescription": "<p>Python, <b>Django</b></p>", "location": "Chennai, India"}},
    }

    def fake(url, data=None, **_):
        if "myworkdayjobs" in url:
            return responses["workday_list"] if url.endswith("/jobs") else responses["workday_detail"]
        if "greenhouse" in url:
            return {"jobs": [{"id": 5, "title": "Backend Engineer", "location": {"name": "Remote"}, "absolute_url": "https://gh/5",
                              "content": "&lt;p&gt;Go&lt;/p&gt;", "updated_at": "2026-10-01"}]}
        if "lever" in url:
            return [{"id": "L1", "text": "Software Engineer", "categories": {"location": "Bangalore"}, "hostedUrl": "https://lv/L1",
                     "descriptionPlain": "Python", "lists": [{"text": "Requirements", "content": "<li>SQL</li>"}], "createdAt": 1760000000000}]
        if "smartrecruiters" in url and "/postings/" in url:
            return {"jobAd": {"sections": {"qualifications": {"text": "<p>Python</p>"}}}}
        if "smartrecruiters" in url:
            return {"content": [{"id": "S1", "name": "Software Engineer", "location": {"city": "Bangalore", "country": "in"},
                                 "releasedDate": "2026-10-10"}]}
        if "amazon.jobs" in url:
            return {"jobs": [{"id_icims": "A1", "title": "Software Development Engineer", "normalized_location": "Chennai, IND",
                              "job_path": "/en/jobs/A1", "description": "x", "basic_qualifications": "Python", "preferred_qualifications": ""}]}
        if "ashbyhq" in url:
            return {"jobs": [{"id": "B1", "title": "Software Engineer", "location": "Remote", "jobUrl": "https://ab/B1", "descriptionPlain": "Python"}]}
        raise AssertionError(url)

    for mod in (workday, greenhouse, lever, smartrecruiters, amazon, ashby):
        monkeypatch.setattr(mod, "http_json", fake)
    kw = ["software", "backend"]
    w = workday.fetch("ADOBE", {"host": "adobe.wd5.myworkdayjobs.com", "tenant": "adobe", "site": "ext"}, kw)
    assert len(w) == 1 and w[0]["description"] == "Python, Django" and w[0]["url"].endswith("/ext/job/Chennai/SE_R1")
    assert greenhouse.fetch("X", {"board": "x"}, kw)[0]["description"] == "Go"
    assert "SQL" in lever.fetch("X", {"site": "x"}, kw)[0]["description"]
    assert smartrecruiters.fetch("VISA", {"company_id": "Visa"}, kw)[0]["description"] == "Python"
    assert amazon.fetch("AMAZON", {}, ["software"])[0]["url"] == "https://www.amazon.jobs/en/jobs/A1"
    assert ashby.fetch("X", {"board": "x"}, kw)[0]["external_id"] == "B1"


def test_discover_queue_and_dedupe(monkeypatch):
    monkeypatch.setitem(sources.ADAPTERS, "adzuna", lambda co, params, kw: [
        job(company=co, title="Python Developer", url=f"https://{co}/1", description="Python Django SQL Docker REST APIs Git"),
        job(company=co, title="Senior Staff Python Developer", url=f"https://{co}/2")])
    from careerpilot import companies
    monkeypatch.setattr(companies, "load", lambda: [{"name": "TCS", "tiers": ["DAY SHARING"], "ats": "adzuna", "params": {}}])
    s1 = discover.run(log=lambda *_: None)
    assert s1["new"] == 2 and s1["queued"] == 1
    s2 = discover.run(log=lambda *_: None)
    assert s2["new"] == 0 and s2["duplicates"] == 2
    pend = apply.pending()
    assert len(pend) == 1 and pend[0]["payload"]["fields"]["email"] == "t@example.com"
    with pytest.raises(PermissionError):
        apply.assist(pend[0]["id"])           # not approved yet → refused
    apply.decide(pend[0]["id"], True)
    apply.mark_submitted(pend[0]["id"])
    assert db.already_applied(pend[0]["job_id"])
    assert apply.queue(pend[0]["job_id"]) is None  # never re-queue an applied job


def test_reports(monkeypatch):
    j1, _ = db.upsert_job(job(salary_text="6 - 9 LPA")); evaluate.evaluate(j1)
    j2, _ = db.upsert_job(job(company="ZOHO", title="Backend Engineer", url="https://z")); evaluate.evaluate(j2)
    db.set_status(j1, "Applied"); db.set_status(j1, "Interview", source="gmail")
    db.set_status(j2, "Applied"); db.set_status(j2, "Rejected")
    path = reports.excel("weekly")
    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert {"Summary", "Top Skills", "Companies", "Skill Gaps & Learning", "Jobs", "Applications", "Funnel", "Trends"} <= set(wb.sheetnames)
    funnel = {r[0].value: r[1].value for r in wb["Funnel"].iter_rows(min_row=2)}
    assert funnel["Applied"] == 2 and funnel["Interview+"] == 1 and funnel["Rejected"] == 1
    out = reports.powerbi()
    assert (out / "fact_jobs.csv").read_text().count("\n") == 3


def test_api_tools_and_guardrails():
    jid, _ = db.upsert_job(job()); evaluate.evaluate(jid)
    assert api.call("search_jobs", {"query": "Python"})[0]["id"] == jid
    assert "PostgreSQL" in api.call("job_details", {"job_id": jid})["gaps"]
    assert "error" in api.call("set_preference", {"section": "automation", "key": "auto_submit_simple_forms", "value": True})
    assert api.call("set_preference", {"section": "preferences", "key": "max_years_required", "value": 3})["ok"]
    assert "approve" not in {t["name"] for t in api.TOOLS}
    assert all(t["input_schema"]["additionalProperties"] is False for t in api.TOOLS)
    assert api.call("update_status", {"job_id": jid, "status": "Bogus"})["error"].startswith("ValueError")
