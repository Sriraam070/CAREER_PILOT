"""Job-market and pipeline reports → Excel (with charts) and Power BI-ready CSVs."""
import csv
import datetime as dt
import json
import re
from collections import Counter, defaultdict

from jhb import skills as sk

from . import db
from .config import REPORTS, cv_text
from .evaluate import DISCLAIMER, parse_salary_lpa

LEARN = {  # skill → (why/where to learn, proof project)
    "FastAPI": ("Official FastAPI tutorial (fastapi.tiangolo.com) → 2 days", "Async CRUD API + JWT + tests"),
    "PostgreSQL": ("PostgreSQL Tutorial (postgresqltutorial.com); port a MySQL project", "Move your Django app to Postgres"),
    "Django REST Framework": ("DRF official tutorial (django-rest-framework.org)", "REST API with viewsets, auth, pagination"),
    "Redis": ("Redis University RU101 (free)", "Cache a slow endpoint and measure it"),
    "Celery": ("Celery 'First steps' + Django integration docs", "Async PDF/report generation"),
    "Docker": ("Docker 'Get started' + Compose", "Dockerize web + db + redis"),
    "Kubernetes": ("Kubernetes basics (kubernetes.io) + minikube", "Deploy your app to minikube"),
    "AWS": ("AWS Cloud Practitioner (Skill Builder, free)", "Deploy to EC2/ECS + S3"),
    "CI/CD": ("GitHub Actions docs", "Lint + test + build pipeline"),
    "React": ("react.dev Learn section", "Frontend for your API"),
    "JavaScript": ("javascript.info", "Small interactive frontend"),
    "TypeScript": ("TypeScript handbook", "Typed frontend"),
    "Java": ("Java MOOC (java-programming.mooc.fi)", "Spring Boot REST API"),
    "Spring": ("Spring Boot guides (spring.io/guides)", "Spring Boot REST API"),
    "System Design": ("'System Design Primer' (GitHub) + ByteByteGo", "Write up the design of your QC app"),
    "Data Structures & Algorithms": ("NeetCode 150 / Striver SDE sheet", "1–2 problems a day"),
    "Microservices": ("microservices.io patterns", "Split one service out of your app"),
    "Kafka": ("Confluent Kafka 101 (free)", "Event-driven notification service"),
    "LLMs": ("Anthropic/Claude docs + DeepLearning.AI short courses", "LLM feature in your project"),
    "Machine Learning": ("Andrew Ng ML Specialization", "Revisit your CNN project with a write-up"),
    "Testing": ("pytest docs + 'Test-Driven Development with Python'", "Raise coverage to 80%+"),
    "Linux": ("Linux Journey (linuxjourney.com)", "Deploy without a PaaS"),
    "Git": ("Pro Git book (git-scm.com/book)", "Clean PR-based history on GitHub"),
    "GCP": ("Google Cloud Skills Boost", "Deploy to Cloud Run"),
    "Azure": ("Microsoft Learn AZ-900 path", "Deploy to App Service"),
}

ROLE_BUCKETS = [
    ("Data / ML", r"data|machine learning|\bml\b|\bai\b|analytics|scientist"),
    ("DevOps / Cloud / SRE", r"devops|sre|site reliability|cloud|platform|infrastructure"),
    ("Frontend", r"front ?end|ui developer|react|angular"),
    ("Full Stack", r"full ?stack"),
    ("Backend", r"back ?end|python|django|java|golang|api"),
    ("QA / Test", r"\bqa\b|test|quality"),
    ("Software Engineer (general)", r"software|developer|engineer|programmer|sde"),
]


def role_bucket(title):
    t = title.lower()
    return next((b for b, rx in ROLE_BUCKETS if re.search(rx, t)), "Other")


def period_range(period="weekly", start=None, end=None):
    today = dt.date.today()
    if period == "custom":
        return dt.date.fromisoformat(start), dt.date.fromisoformat(end)
    days = {"daily": 1, "weekly": 7, "monthly": 30, "all": 3650}[period]
    return today - dt.timedelta(days=days - 1), today


def collect(start, end):
    s, e = start.isoformat(), (end + dt.timedelta(days=1)).isoformat()
    jobs = db.rows("SELECT j.*, c.tiers FROM jobs j LEFT JOIN companies c ON c.name=j.company "
                   "WHERE j.discovered_at >= ? AND j.discovered_at < ?", (s, e))
    apps = db.rows("""SELECT a.*, j.company, j.title, j.url, j.match_score, j.probability, j.location
                      FROM applications a JOIN jobs j ON j.id=a.job_id""")
    hist = db.rows("""SELECT h.*, j.company, j.title FROM status_history h JOIN applications a ON a.id=h.application_id
                      JOIN jobs j ON j.id=a.job_id WHERE h.at >= ? AND h.at < ?""", (s, e))
    for j in jobs:
        j["skills_l"] = json.loads(j["skills"] or "[]")
        j["gaps_l"] = json.loads(j["gaps"] or "[]")
        j["tiers_l"] = json.loads(j["tiers"] or "[]")
        j["role"] = role_bucket(j["title"])
        j["salary"] = parse_salary_lpa(j["salary_text"] or j["description"] or "")
    return jobs, apps, hist


def metrics(jobs, apps, hist):
    n = len(jobs)
    skill_count = Counter(s for j in jobs for s in j["skills_l"])
    _, cv_all = sk._split_cv(cv_text())
    skills_sec = sk._split_cv(cv_text())[0]

    def have(skill):
        pats = sk._COMPILED[skill][1]
        return "Yes" if any(p.search(skills_sec) for p in pats) else ("In CV text" if any(p.search(cv_all) for p in pats) else "No")

    top = [(s, sk.VOCAB[s][0], c, round(100 * c / n, 1) if n else 0, have(s)) for s, c in skill_count.most_common(40)]
    gaps = [(s, cat, c, pct) for s, cat, c, pct, h in top if h == "No" and cat != "Soft Skills"]
    by_co = defaultdict(list)
    for j in jobs:
        by_co[j["company"]].append(j)
    stage = Counter(a["status"] for a in apps)
    applied = sum(v for k, v in stage.items() if k != "Not Applied")
    reached = lambda *st: sum(stage[s] for s in st)
    funnel = [("Applied", applied), ("OA/Test+", reached("OA/Test", "Interview", "Offer")),
              ("Interview+", reached("Interview", "Offer")), ("Offer", stage["Offer"]), ("Rejected", stage["Rejected"])]
    conv = {"response_rate": round(100 * reached("OA/Test", "Interview", "Offer") / applied, 1) if applied else 0,
            "interview_rate": round(100 * reached("Interview", "Offer") / applied, 1) if applied else 0,
            "offer_rate": round(100 * stage["Offer"] / applied, 1) if applied else 0,
            "rejection_rate": round(100 * stage["Rejected"] / applied, 1) if applied else 0}
    return {"n_jobs": n, "n_companies": len(by_co), "top": top, "gaps": gaps, "by_co": by_co,
            "stage": stage, "funnel": funnel, "conv": conv,
            "avg_match": round(sum(j["match_score"] or 0 for j in jobs) / n, 1) if n else 0}


def excel(period="weekly", start=None, end=None):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    start, end = period_range(period, start, end)
    jobs, apps, hist = collect(start, end)
    m = metrics(jobs, apps, hist)
    wb = Workbook()
    H = PatternFill("solid", fgColor="1F4E78")

    def sheet(title, header, data, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), H
        for r in data:
            ws.append(list(r))
        for i, w in enumerate(widths or [18] * len(header), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if data:
            ws.auto_filter.ref = ws.dimensions
        return ws

    def bar(ws, title, col, n, anchor, horizontal=True):
        if n == 0:
            return
        ch = BarChart(); ch.type = "bar" if horizontal else "col"; ch.title = title; ch.legend = None
        ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        ch.height, ch.width = 9 + n * 0.25, 18
        ws.add_chart(ch, anchor)

    # Summary
    ws = wb.active; ws.title = "Summary"
    ws["A1"] = f"CareerPilot report: {period} ({start} → {end})"; ws["A1"].font = Font(bold=True, size=14)
    kpis = [("Jobs discovered", m["n_jobs"]), ("Companies hiring", m["n_companies"]),
            ("Average skill match %", m["avg_match"]), ("Applications (all time)", sum(m["stage"].values())),
            ("Response rate % (OA/interview/offer ÷ applied)", m["conv"]["response_rate"]),
            ("Interview conversion %", m["conv"]["interview_rate"]), ("Offer rate %", m["conv"]["offer_rate"]),
            ("Rejection rate %", m["conv"]["rejection_rate"]), ("Status changes in period", len(hist))]
    for i, (k, v) in enumerate(kpis, 3):
        ws.cell(i, 1, k).font = Font(bold=True); ws.cell(i, 2, v)
    ws.cell(len(kpis) + 4, 1, "Top 5 skills to learn next: " + ", ".join(g[0] for g in m["gaps"][:5]))
    ws.cell(len(kpis) + 5, 1, f"* {DISCLAIMER}").font = Font(italic=True, color="888888")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 48, 14

    ws = sheet("Top Skills", ["Skill", "Category", "Jobs asking", "% of jobs", "In your CV?"], m["top"], [26, 20, 12, 10, 12])
    bar(ws, "Most-requested skills", 3, min(15, len(m["top"])), "G2")

    rows_co = []
    for co, js in sorted(m["by_co"].items(), key=lambda kv: -len(kv[1])):
        yrs = [j["years_required"] for j in js if j["years_required"]]
        sal = [j["salary"] for j in js if j["salary"]]
        rows_co.append((co, ", ".join(js[0]["tiers_l"]), len(js), len({j["role"] for j in js}),
                        round(sum(j["match_score"] or 0 for j in js) / len(js)),
                        f"{min(yrs)}–{max(yrs)}" if yrs else "—",
                        f"{min(s[0] for s in sal)}–{max(s[1] for s in sal)} LPA" if sal else "not stated",
                        "; ".join(sorted({(j["location"] or "—")[:30] for j in js})[:4])))
    ws = sheet("Companies", ["Company", "Tier", "Open positions", "Role types", "Avg match %", "Years asked",
                             "Salary", "Locations"], rows_co, [22, 16, 14, 11, 12, 12, 16, 50])
    bar(ws, "Open positions by company", 3, min(15, len(rows_co)), "J2")

    roles = Counter(j["role"] for j in jobs)
    tech_by_role = defaultdict(Counter)
    for j in jobs:
        tech_by_role[j["role"]].update(s for s in j["skills_l"] if sk.VOCAB[s][0] != "Soft Skills")
    ws = sheet("Roles", ["Role", "Positions", "Top technologies"],
               [(r, c, ", ".join(t for t, _ in tech_by_role[r].most_common(6))) for r, c in roles.most_common()], [28, 11, 70])
    bar(ws, "Positions by role", 2, len(roles), "E2")

    ws = sheet("Skill Gaps & Learning", ["Priority", "Skill", "Category", "Jobs asking", "% of jobs", "How to learn", "Proof project"],
               [(i, s, cat, c, pct, *LEARN.get(s, ("Official docs + a small project", "Add it to a portfolio project")))
                for i, (s, cat, c, pct) in enumerate(m["gaps"][:20], 1)], [8, 24, 18, 11, 10, 55, 40])

    ws = sheet("Jobs", ["Company", "Title", "Role", "Location", "Posted", "Discovered", "Match %", "Est. response %*",
                        "Years", "Salary", "Gaps", "State", "Source", "Link"],
               [(j["company"], j["title"], j["role"], j["location"], (j["posted_at"] or "")[:10], j["discovered_at"][:10],
                 j["match_score"], j["probability"], j["years_required"], j["salary_text"] or "", ", ".join(j["gaps_l"][:6]),
                 j["state"], j["source"], j["url"]) for j in sorted(jobs, key=lambda j: -(j["match_score"] or 0))],
               [18, 38, 16, 22, 11, 11, 9, 12, 7, 14, 40, 10, 12, 50])
    ws.cell(ws.max_row + 2, 1, f"* {DISCLAIMER}").font = Font(italic=True, color="888888")

    ws = sheet("Applications", ["Company", "Title", "Status", "Applied", "Updated", "Next step", "Interview", "Match %", "Link"],
               [(a["company"], a["title"], a["status"], (a["applied_at"] or "")[:10], (a["updated_at"] or "")[:10],
                 a["next_step"] or "", a["interview_at"] or "", a["match_score"], a["url"]) for a in apps],
               [18, 38, 12, 11, 11, 40, 18, 9, 50])

    ws = sheet("Funnel", ["Stage", "Count"], m["funnel"], [16, 10])
    bar(ws, "Application funnel", 2, len(m["funnel"]), "D2", horizontal=False)

    weekly = defaultdict(Counter)
    for h in hist:
        wk = dt.date.fromisoformat(h["at"][:10]).strftime("%Y-W%W")
        weekly[wk][h["status"]] += 1
    st = ["Applied", "OA/Test", "Interview", "Rejected", "Offer"]
    ws = sheet("Trends", ["Week"] + st, [(w, *[weekly[w][s] for s in st]) for w in sorted(weekly)], [12] + [11] * 5)
    if weekly:
        ch = LineChart(); ch.title = "Status changes per week"
        ch.add_data(Reference(ws, min_col=2, max_col=6, min_row=1, max_row=len(weekly) + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(weekly) + 1))
        ws.add_chart(ch, "H2")
    rej = Counter(h["company"] for h in hist if h["status"] == "Rejected")
    ws = sheet("Rejections", ["Company", "Rejections"], rej.most_common(), [24, 12])

    disc = Counter(j["discovered_at"][:10] for j in jobs)
    sheet("Daily Discovery", ["Date", "New jobs"], sorted(disc.items()), [12, 10])

    for w in wb.worksheets:
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
    REPORTS.mkdir(exist_ok=True)
    path = REPORTS / f"careerpilot-{period}-{start}_{end}.xlsx"
    wb.save(path)
    db.audit("report_excel", path.name)
    return path


def powerbi():
    """Star-schema CSVs for Power BI: Get Data → Folder → reports/powerbi."""
    out = REPORTS / "powerbi"
    out.mkdir(parents=True, exist_ok=True)
    jobs, apps, hist = collect(dt.date(2000, 1, 1), dt.date.today())

    def write(name, header, data):
        with open(out / name, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f); w.writerow(header); w.writerows(data)

    write("fact_jobs.csv", ["job_id", "company", "title", "role", "location", "posted_at", "discovered_date", "match_pct",
                            "est_response_pct", "years_required", "salary_min_lpa", "salary_max_lpa", "state", "source", "url"],
          [(j["id"], j["company"], j["title"], j["role"], j["location"], (j["posted_at"] or "")[:10], j["discovered_at"][:10],
            j["match_score"], j["probability"], j["years_required"], *(j["salary"] or ("", "")), j["state"], j["source"], j["url"]) for j in jobs])
    write("bridge_job_skill.csv", ["job_id", "skill", "category", "is_gap"],
          [(j["id"], s, sk.VOCAB[s][0], int(s in j["gaps_l"])) for j in jobs for s in j["skills_l"]])
    write("fact_applications.csv", ["application_id", "job_id", "company", "title", "status", "applied_date", "updated_date", "interview_at"],
          [(a["id"], a["job_id"], a["company"], a["title"], a["status"], (a["applied_at"] or "")[:10], (a["updated_at"] or "")[:10], a["interview_at"]) for a in apps])
    write("fact_status_history.csv", ["application_id", "company", "status", "date", "source"],
          [(h["application_id"], h["company"], h["status"], h["at"][:10], h["source"]) for h in hist])
    cos = db.rows("SELECT name, tiers, ats FROM companies")
    write("dim_company.csv", ["company", "tier", "ats"], [(c["name"], ", ".join(json.loads(c["tiers"] or "[]")), c["ats"]) for c in cos])
    db.audit("report_powerbi", str(out))
    return out
